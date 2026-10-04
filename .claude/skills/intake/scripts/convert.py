#!/usr/bin/env python3
"""Chuyển 1 tài liệu gốc (docx/doc/pdf/xlsx/xls/csv/pptx/ảnh) thành bản nháp Markdown + ảnh.

Dùng:  python3 convert.py <file gốc> --out <thư mục staging> --slug <slug> [--date YYYY-MM-DD]

Ghi ra <out>/:
  {slug}.md          bản nháp Markdown (Claude rà + sửa tiếp)
  images/            ảnh nhúng, ảnh trang scan, ảnh xem trước trang có hình vẽ
  source/{slug}.ext  bản sao file gốc
  _raw.txt           toàn bộ chữ trích thô từ file gốc (để verify.py đối chiếu độ đủ)
  _report.json       thống kê + cảnh báo (hình vẽ/SmartArt/biểu đồ/trang scan/sheet ẩn...)

Không cần mạng. Thư viện: pymupdf, pymupdf4llm, mammoth, markdownify, openpyxl, python-pptx, xlrd
(thiếu → in lệnh pip cài rồi thoát mã 3). OCR dùng Vision của macOS qua ocr.swift (tiếng Việt).
"""
import argparse, csv, datetime as dt, glob, json, os, re, shutil, subprocess, sys, unicodedata, zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
IMG_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".tif", ".tiff", ".bmp", ".heic"}
NEED = {"pymupdf": "pymupdf", "pymupdf4llm": "pymupdf4llm", "mammoth": "mammoth",
        "markdownify": "markdownify", "openpyxl": "openpyxl", "pptx": "python-pptx", "xlrd": "xlrd"}

nfc = lambda s: unicodedata.normalize("NFC", s or "")


def need(*mods):
    import importlib
    miss = [NEED[m] for m in mods if importlib.util.find_spec(m) is None]
    if miss:
        print(f"THIẾU THƯ VIỆN: pip3 install {' '.join(miss)}", file=sys.stderr)
        sys.exit(3)


class Ctx:
    def __init__(self, src, out, slug, img_prefix="", depth=0):
        self.src, self.out, self.slug, self.img_prefix, self.depth = src, out, slug, img_prefix, depth
        self.attachments = []  # (đường dẫn tương đối trong source/, tên hiển thị)
        self.img_dir = os.path.join(out, "images")
        os.makedirs(self.img_dir, exist_ok=True)
        self.raw, self.warn, self.stats, self.ocr_queue = [], [], {}, []
        self.raw_ocr = []  # chữ OCR (Claude sẽ sửa theo ảnh → verify chỉ báo tham khảo)
        self.n_img = 0

    def save_img(self, data, ext, name=None):
        self.n_img += 1
        ext = ext.lower().lstrip(".").replace("jpeg", "jpg")
        name = name or f"{self.img_prefix}img-{self.n_img:02d}.{ext}"
        with open(os.path.join(self.img_dir, name), "wb") as f:
            f.write(data)
        return f"images/{name}"


def esc_cell(v):
    s = "" if v is None else str(v)
    return nfc(s).replace("|", "\\|").replace("\r", "").replace("\n", "<br>").strip()


def md_table(rows):
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    rows = [list(r) + [""] * (w - len(r)) for r in rows]
    out = ["| " + " | ".join(esc_cell(c) for c in rows[0]) + " |", "|" + "---|" * w]
    out += ["| " + " | ".join(esc_cell(c) for c in r) + " |" for r in rows[1:]]
    return "\n".join(out)


def fix_md(md, demote=1):
    """Hạ cấp heading (H1 dành cho tiêu đề tài liệu) + đẩy hàng đầu lên làm tiêu đề bảng khi tiêu đề trống."""
    md = re.sub(r"^(#{1,6}) (.*)$", lambda m: "#" * min(6, len(m.group(1)) + demote) + " " +
                re.sub(r"\*{2,3}([^*]+)\*{2,3}", r"\1", m.group(2)).strip(), md, flags=re.M)  # bỏ đậm thừa trong heading
    md = re.sub(r"^\|(?:\s*\|)+\n\|(?:\s*-+\s*\|)+\n(\|.*\|)\n",
                lambda m: m.group(1) + "\n|" + "---|" * (m.group(1).count("|") - 1) + "\n", md, flags=re.M)
    return re.sub(r"\n{3,}", "\n\n", md).strip()


def run_ocr(ctx):
    """OCR tất cả ảnh trong hàng đợi bằng 1 lần gọi swift. Trả {đường dẫn tương đối: chữ}."""
    if not ctx.ocr_queue:
        return {}
    paths = [os.path.join(ctx.out, p) for p in ctx.ocr_queue]
    try:
        res = subprocess.run(["swift", os.path.join(HERE, "ocr.swift"), *paths],
                             capture_output=True, text=True, timeout=900)
    except Exception as e:  # không có swift (không phải macOS)
        ctx.warn.append(f"OCR không chạy được ({e}) — Claude phải đọc ảnh trực tiếp để gõ lại chữ")
        return {}
    out, cur = {}, None
    for ln in res.stdout.splitlines():
        if ln.startswith("=== "):
            cur = os.path.relpath(ln[4:], ctx.out)
            out[cur] = []
        elif cur:
            out[cur].append(ln)
    return {k: nfc("\n".join(v)).strip() for k, v in out.items()}


# ---------------- FILE NHÚNG (OLE / Office trong Office) ----------------
def extract_embedding(data, name):
    """Trả (bytes, đuôi, tên gốc|None) của file thật nằm trong 1 file nhúng; None nếu không nhận ra."""
    ext = os.path.splitext(name)[1].lower()
    if ext in (".xlsx", ".xlsm", ".docx", ".pptx", ".pdf"):
        return data, ext, None
    try:
        import olefile, io, struct
        ole = olefile.OleFileIO(io.BytesIO(data))
        streams = ["/".join(x) for x in ole.listdir()]
        if "CONTENTS" in streams:                     # PDF nhúng bằng Acrobat
            d = ole.openstream("CONTENTS").read()
            if d.startswith(b"%PDF"):
                return d, ".pdf", None
        if "\x01Ole10Native" in streams:              # "Package" — file bất kỳ kèm tên gốc
            d = ole.openstream("\x01Ole10Native").read()
            i = 6
            fn_end = d.index(b"\0", i); fn = d[i:fn_end].decode("latin-1", "ignore"); i = fn_end + 1
            i = d.index(b"\0", i) + 1 + 8               # đường dẫn gốc + 8 byte
            i = d.index(b"\0", i) + 1                   # đường dẫn tạm
            n = struct.unpack("<I", d[i:i + 4])[0]
            return d[i + 4:i + 4 + n], os.path.splitext(fn)[1].lower() or ".bin", fn
        if "Package" in streams:                       # Office nhúng dạng OLE
            d = ole.openstream("Package").read()
            kind = b"xl/" in d[:4000] and ".xlsx" or b"word/" in d[:4000] and ".docx" or ".pptx"
            return d, kind, None
    except Exception:
        pass
    i, j = data.find(b"%PDF"), data.rfind(b"%%EOF")
    if i >= 0 and j > i:
        return data[i:j + 5], ".pdf", None
    return None


def attach(ctx, data, ext, label=None):
    """Lưu file đính kèm vào source/ và ghi nhận để chuyển nội dung ở cuối tài liệu."""
    n = len(ctx.attachments) + 1
    rel = f"source/{ctx.slug}-dinh-kem-{n:02d}{ext}"
    os.makedirs(os.path.join(ctx.out, "source"), exist_ok=True)
    with open(os.path.join(ctx.out, rel), "wb") as f:
        f.write(data)
    ctx.attachments.append((rel, label or f"Đính kèm {n:02d}{ext}"))
    return n, rel


def find_duplicate(path, md_body, root):
    """Đính kèm đã có trong docs-input chưa? Trả (đường dẫn, 'y hệt'|'gần như trùng NN%') hoặc None."""
    if not root or not os.path.isdir(root):
        return None
    import hashlib
    h = hashlib.sha1(open(path, "rb").read()).hexdigest()
    for dp, dns, fns in os.walk(root):
        dns[:] = [d for d in dns if not d.startswith(".")]
        for fn in fns:
            fp = os.path.join(dp, fn)
            if fn.startswith(".") or os.path.getsize(fp) != os.path.getsize(path):
                continue
            if hashlib.sha1(open(fp, "rb").read()).hexdigest() == h:
                if os.path.basename(dp) == "source":  # bản gốc của 1 tài liệu đã chuyển → trỏ .md
                    doc = os.path.dirname(dp)
                    md = os.path.join(doc, os.path.basename(doc) + ".md")
                    fp = md if os.path.exists(md) else fp
                return os.path.relpath(fp, os.path.dirname(root)), "y hệt"
    sys.path.insert(0, HERE)
    from match import body_words, shingles
    mine = shingles(body_words(md_body))
    best = (0, None)
    for fp in glob.glob(os.path.join(root, "**/*.md"), recursive=True):
        rel = os.path.relpath(fp, root)
        if rel == "README.md" or rel.startswith(("_", ".")):
            continue
        other = shingles(body_words(open(fp, encoding="utf-8").read()))
        if mine and other:
            c = len(mine & other) / min(len(mine), len(other))
            if c > best[0]:
                best = (c, fp)
    if best[0] >= 0.6:
        return os.path.relpath(best[1], os.path.dirname(root)), f"gần như trùng {best[0]:.0%}"
    return None


# ---------------- DOCX ----------------
def conv_docx(ctx):
    need("mammoth", "markdownify")
    import mammoth
    from markdownify import markdownify

    # icon của file nhúng (EMF trong <w:object>) → file nhúng tương ứng
    z0 = zipfile.ZipFile(ctx.src)
    rels = z0.read("word/_rels/document.xml.rels").decode("utf8", "ignore") if "word/_rels/document.xml.rels" in z0.namelist() else ""
    rid = dict(re.findall(r'Id="([^"]+)"[^>]*?Target="([^"]+)"', rels))
    rid.update({k: v for v, k in re.findall(r'Target="([^"]+)"[^>]*?Id="([^"]+)"', rels)})
    media = {z0.read("word/" + t): t for t in rid.values() if t.startswith("media/") and "word/" + t in z0.namelist()}
    icon_to_emb = {}
    for obj in re.findall(r"<w:object\b.*?</w:object>", z0.read("word/document.xml").decode("utf8", "ignore"), re.S):
        im = re.search(r'<v:imagedata[^>]*r:id="([^"]+)"', obj)
        em = re.search(r'<o:OLEObject[^>]*r:id="([^"]+)"', obj)
        if im and em and im.group(1) in rid and em.group(1) in rid:
            icon_to_emb[rid[im.group(1)]] = rid[em.group(1)]
    used_emb = set()

    def img_handler(image):
        with image.open() as f:
            data = f.read()
        emb = icon_to_emb.get(media.get(data))
        if emb and ctx.depth < 2:
            used_emb.add(emb)
            got = extract_embedding(z0.read("word/" + emb), emb)
            if got:
                n, rel = attach(ctx, got[0], got[1], got[2])
                return {"src": rel, "alt": f"DINHKEM{n:02d}"}
            ctx.warn.append(f"File nhúng {emb} không tách được — mở bản gốc để xem")
        ext = (image.content_type or "image/png").split("/")[-1].replace("x-", "")
        if ext in ("emf", "wmf", "x-emf", "x-wmf"):
            ctx.warn.append(f"Ảnh dạng {ext.upper()} (thường là sơ đồ vẽ trong Word) — không hiển thị được trong .md; "
                            "cần dựng lại bằng mermaid hoặc xin ảnh chụp")
        return {"src": ctx.save_img(data, ext)}

    sys.path.insert(0, HERE)
    from docx_numbering import number_docx
    numbered = os.path.join(ctx.out, "_numbered.docx")
    try:
        n_num, n_head = number_docx(ctx.src, numbered)
    except Exception as e:
        numbered, n_num, n_head = ctx.src, 0, 0
        ctx.warn.append(f"Không dựng lại được số thứ tự tự động ({e}) — số mục 1./2.2/a) có thể bị mất, đối chiếu bản gốc")
    smap = "p[style-name='Title'] => h1:fresh\np[style-name='Subtitle'] => h2:fresh\n" + \
        "\n".join(f"p.Heading{i} => h{i}:fresh" for i in range(1, 7))
    with open(numbered, "rb") as f:
        res = mammoth.convert_to_html(f, convert_image=mammoth.images.img_element(img_handler), style_map=smap)
    if numbered != ctx.src:
        os.remove(numbered)
    if n_num:
        ctx.stats.update(numbered_paragraphs=n_num, numbered_headings=n_head)
    for m in res.messages:
        if "image" in str(m.message).lower():
            ctx.warn.append(f"mammoth: {m.message}")
    md = fix_md(markdownify(res.value, heading_style="ATX", bullets="-", strip=["span"]))
    md = re.sub(r"!\[DINHKEM(\d+)\]\(([^)]+)\)",
                lambda m: f"📎 **Đính kèm {m.group(1)}:** [{os.path.basename(m.group(2))}]({m.group(2)}) "
                          f"— nội dung ở mục *Đính kèm {m.group(1)}* cuối tài liệu", md)
    for t in rid.values():  # file nhúng không có icon trong thân bài
        if t.startswith("embeddings/") and t not in used_emb and ctx.depth < 2:
            got = extract_embedding(z0.read("word/" + t), t)
            if got:
                attach(ctx, got[0], got[1], got[2])

    # trích thô toàn bộ chữ (kể cả text box, header/footer, chú thích) + dò thứ không chuyển được
    z = zipfile.ZipFile(ctx.src)
    names = z.namelist()
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    import xml.etree.ElementTree as ET

    def paras(xml):
        root = ET.fromstring(xml)
        for p in root.iter(W + "p"):
            st = p.find(f"{W}pPr/{W}pStyle")
            if st is not None and st.get(W + "val", "").upper().startswith("TOC"):
                continue  # dòng mục lục (kèm số trang) — không phải nội dung
            t = "".join(x.text or "" for x in p.iter(W + "t"))
            if t.strip():
                yield t
    doc = z.read("word/document.xml")
    ctx.raw += list(paras(doc))
    extra = [n for n in names if re.match(r"word/(header|footer|footnotes|endnotes|comments)\d*\.xml$", n)]
    hf = []
    for n in extra:
        hf += [f"[{os.path.basename(n)[:-4]}] {t}" for t in paras(z.read(n))]
    if hf:
        ctx.raw += hf
        ctx.warn.append(f"Có {len(hf)} đoạn trong header/footer/chú thích/comment — mammoth bỏ qua, "
                        "Claude quyết định giữ (thường giữ chú thích + comment, bỏ header/footer lặp)")
    s = doc.decode("utf8", "ignore")
    shapes = s.count("<wps:wsp")
    txbx = s.count("<w:txbxContent")
    smart = len([n for n in names if n.startswith("word/diagrams/")])
    charts = len([n for n in names if re.match(r"word/charts/chart\d+\.xml$", n)])
    embeds = len([n for n in names if n.startswith("word/embeddings/")])
    tracked = s.count("<w:ins ") + s.count("<w:del ")
    if shapes or txbx:
        ctx.warn.append(f"Có {shapes} hình vẽ (shape) / {txbx} text box trong Word — chữ trong đó có ở _raw.txt "
                        "nhưng KHÔNG có trong bản nháp; dựng lại sơ đồ bằng mermaid hoặc ghi lại nội dung")
    if smart:
        ctx.warn.append("Có SmartArt — không chuyển được tự động, dựng lại bằng mermaid/danh sách")
    if charts:
        ctx.warn.append(f"Có {charts} biểu đồ Excel nhúng — ghi lại số liệu dạng bảng")
    if embeds and len(ctx.attachments) < embeds:
        ctx.warn.append(f"Có {embeds} file nhúng, mới tách được {len(ctx.attachments)} — mở bản gốc kiểm phần còn lại")
    if tracked:
        ctx.warn.append(f"Có {tracked} chỗ Track Changes chưa chấp nhận — hỏi user lấy bản nào")
    title = None
    try:
        core = z.read("docProps/core.xml").decode("utf8", "ignore")
        m = re.search(r"<dc:title>([^<]+)</dc:title>", core)
        title = m.group(1).strip() if m else None
    except KeyError:
        pass
    ctx.stats.update(paragraphs=len(ctx.raw), images=ctx.n_img, tables=res.value.count("<table"),
                     shapes=shapes, smartart=bool(smart), charts=charts)
    return title, md


# ---------------- PDF ----------------
def conv_pdf(ctx):
    need("pymupdf", "pymupdf4llm")
    import pymupdf, pymupdf4llm
    doc = pymupdf.open(ctx.src)
    title = (doc.metadata or {}).get("title") or None
    scanned, previews = [], []
    for i, page in enumerate(doc, 1):
        t = page.get_text()
        if len(t.strip()) < 30 and page.get_images():
            scanned.append(i)
        elif len(page.get_drawings()) > 40 or len(page.get_images()) > 0:
            previews.append(i)
    try:
        chunks = pymupdf4llm.to_markdown(ctx.src, page_chunks=True, write_images=True,
                                         image_path=ctx.img_dir, image_format="png", dpi=150,
                                         show_progress=False)
    except Exception as e:
        ctx.warn.append(f"pymupdf4llm lỗi ({e}) — dùng chữ thô từng trang")
        chunks = [{"text": p.get_text()} for p in doc]
    parts = []
    for i, page in enumerate(doc, 1):
        body = chunks[i - 1]["text"] if i - 1 < len(chunks) else ""
        body = re.sub(r"\]\(([^)]*/)?images/([^)/]+)\)", r"](images/\2)", body)  # đường dẫn tuyệt đối → tương đối
        if i in scanned:
            pix = page.get_pixmap(dpi=200)
            rel = f"images/{ctx.img_prefix}page-{i:02d}.png"
            pix.save(os.path.join(ctx.out, rel))
            ctx.ocr_queue.append(rel)
            parts.append((i, rel, None))
        else:
            ctx.raw.append(page.get_text())
            prev = None
            if i in previews:
                prev = f"images/{ctx.img_prefix}page-{i:02d}.png"
                page.get_pixmap(dpi=110).save(os.path.join(ctx.out, prev))
            parts.append((i, prev, body.strip()))
    ocr = run_ocr(ctx)
    md = []
    for i, img, body in parts:
        md.append(f"## Trang {i}\n")
        if body is None:  # trang scan
            txt = ocr.get(img, "")
            ctx.raw_ocr.append(txt)
            md.append(f"![Trang {i} — bản scan]({img})\n")
            md.append("<!-- OCR thô — Claude đối chiếu ảnh để sửa dấu/bảng -->\n" + (txt or "_(OCR không ra chữ)_") + "\n")
        else:
            md.append(fix_md(body, 2) + "\n")
            if img:
                md.append(f"![Trang {i} — ảnh xem trước để đối chiếu hình vẽ]({img})\n")
    if scanned:
        ctx.warn.append(f"{len(scanned)} trang scan (trang {', '.join(map(str, scanned))}) — đã OCR; "
                        "Claude PHẢI xem ảnh trang để sửa chữ OCR và dựng lại bảng/biểu mẫu")
    if previews:
        ctx.warn.append(f"Trang có hình vẽ/ảnh: {', '.join(map(str, previews))} — đã lưu ảnh xem trước; "
                        "xem ảnh để ghi lại sơ đồ (mermaid) nếu chữ trong hình không có trong bản nháp")
    ctx.stats.update(pages=len(doc), scanned_pages=scanned, preview_pages=previews,
                     images=len(os.listdir(ctx.img_dir)))
    return title, "\n".join(md)


# ---------------- XLSX / XLS / CSV ----------------
def fmt_val(v, numfmt=""):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float):
        if "%" in (numfmt or ""):
            return f"{v * 100:g}%"
        return f"{v:.10g}"
    return str(v)


def trim(grid):
    grid = [r for r in grid]
    while grid and not any(str(c).strip() for c in grid[-1]):
        grid.pop()
    while grid and not any(str(c).strip() for c in grid[0]):
        grid.pop(0)
    if not grid:
        return grid
    w = max((max((j for j, c in enumerate(r) if str(c).strip()), default=-1) for r in grid), default=-1) + 1
    grid = [r[:w] for r in grid]
    # bỏ cột trống hoàn toàn ở bên trái
    left = min((next((j for j, c in enumerate(r) if str(c).strip()), w) for r in grid), default=0)
    return [r[left:] for r in grid]


def conv_xlsx(ctx):
    need("openpyxl")
    import openpyxl
    wb_v = openpyxl.load_workbook(ctx.src, data_only=True)
    wb_f = openpyxl.load_workbook(ctx.src, data_only=False)
    md, sheets = [], []
    for ws in wb_v.worksheets:
        wf = wb_f[ws.title]
        merged = list(ws.merged_cells.ranges)
        grid = []
        nform = 0
        for r_v, r_f in zip(ws.iter_rows(), wf.iter_rows()):
            row = []
            for c, cf in zip(r_v, r_f):
                if isinstance(cf.value, str) and cf.value.startswith("="):
                    nform += 1
                row.append(fmt_val(c.value, c.number_format))
            grid.append(row)
        grid = trim(grid)
        state = ws.sheet_state
        sheets.append(dict(name=ws.title, rows=len(grid), state=state, merged=len(merged), formulas=nform))
        md.append(f"## Sheet: {ws.title}" + (" _(sheet ẩn)_" if state != "visible" else "") + "\n")
        if not grid:
            md.append("_(sheet trống)_\n")
            continue
        md.append(md_table(grid) + "\n")
        ctx.raw += ["\t".join(r) for r in grid]
        if merged:
            ctx.warn.append(f"Sheet '{ws.title}': {len(merged)} vùng gộp ô — giá trị chỉ nằm ở ô đầu; "
                            "kiểm tra tiêu đề nhiều tầng có đọc đúng không")
        if nform:
            md.append(f"> Sheet có {nform} ô công thức — bảng trên là giá trị đã tính.\n")
        if state != "visible":
            ctx.warn.append(f"Sheet '{ws.title}' bị ẩn — đã chuyển, hỏi user có cần giữ không")
        for img in getattr(ws, "_images", []):
            try:
                rel = ctx.save_img(img._data(), img.format or "png")
                md.append(f"![Ảnh trong sheet {ws.title}]({rel})\n")
            except Exception:
                ctx.warn.append(f"Sheet '{ws.title}' có ảnh không trích được")
        if getattr(ws, "_charts", []):
            ctx.warn.append(f"Sheet '{ws.title}' có {len(ws._charts)} biểu đồ — số liệu nằm trong bảng, "
                            "ghi chú biểu đồ thể hiện gì")
        cmts = [(c.coordinate, c.comment.text) for r in ws.iter_rows() for c in r if c.comment]
        if cmts:
            md.append("**Ghi chú trong ô:**\n\n" + "\n".join(f"- `{a}`: {esc_cell(t)}" for a, t in cmts) + "\n")
            ctx.raw += [t for _, t in cmts]
    ctx.stats.update(sheets=sheets)
    return wb_v.properties.title, "\n".join(md)


def conv_xls(ctx):
    need("xlrd")
    import xlrd
    wb = xlrd.open_workbook(ctx.src)
    md, sheets = [], []
    for sh in wb.sheets():
        grid = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    v = xlrd.xldate.xldate_as_datetime(v, wb.datemode)
                row.append(fmt_val(v))
            grid.append(row)
        grid = trim(grid)
        sheets.append(dict(name=sh.name, rows=len(grid)))
        md.append(f"## Sheet: {sh.name}\n")
        md.append((md_table(grid) if grid else "_(sheet trống)_") + "\n")
        ctx.raw += ["\t".join(r) for r in grid]
    ctx.stats.update(sheets=sheets)
    return None, "\n".join(md)


def conv_csv(ctx):
    with open(ctx.src, encoding="utf-8-sig", errors="replace") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        grid = trim([r for r in csv.reader(f, dialect)])
    ctx.raw += ["\t".join(r) for r in grid]
    ctx.stats.update(rows=len(grid))
    return None, md_table(grid)


# ---------------- PPTX ----------------
def conv_pptx(ctx):
    need("pptx")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(ctx.src)
    md = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text_frame.text if slide.shapes.title is not None else ""
        md.append(f"## Slide {i}" + (f" — {nfc(title)}" if title else "") + "\n")

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from walk(sh.shapes)
                else:
                    yield sh
        for sh in walk(slide.shapes):
            if sh.has_text_frame and sh != slide.shapes.title:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        md.append("  " * p.level + f"- {nfc(t)}")
                        ctx.raw.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text for c in r.cells] for r in sh.table.rows]
                md.append("\n" + md_table(rows) + "\n")
                ctx.raw += ["\t".join(r) for r in rows]
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                rel = ctx.save_img(sh.image.blob, sh.image.ext)
                md.append(f"\n![Ảnh slide {i}]({rel})\n")
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ctx.warn.append(f"Slide {i} có biểu đồ — ghi lại số liệu dạng bảng")
        if title:
            ctx.raw.append(title)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            n = slide.notes_slide.notes_text_frame.text.strip()
            md.append(f"\n> Ghi chú người trình bày: {esc_cell(n)}\n")
            ctx.raw.append(n)
        md.append("")
    ctx.stats.update(slides=len(prs.slides), images=ctx.n_img)
    return prs.core_properties.title, "\n".join(md)


# ---------------- ẢNH ----------------
def conv_image(ctx):
    ext = os.path.splitext(ctx.src)[1].lower()
    name = f"{ctx.slug}{'.jpg' if ext in ('.heic', '.tif', '.tiff', '.bmp') else ext.replace('jpeg', 'jpg')}"
    dst = os.path.join(ctx.img_dir, name)
    if ext in (".heic", ".tif", ".tiff", ".bmp"):
        subprocess.run(["sips", "-s", "format", "jpeg", ctx.src, "--out", dst], capture_output=True)
    else:
        shutil.copy2(ctx.src, dst)
    rel = f"images/{name}"
    ctx.ocr_queue.append(rel)
    txt = run_ocr(ctx).get(rel, "")
    ctx.raw_ocr.append(txt)
    ctx.warn.append("Tài liệu là ảnh — đã OCR; Claude PHẢI xem ảnh để sửa chữ, dựng lại bảng/biểu mẫu "
                    "và mô tả phần không phải chữ (sơ đồ, con dấu, QR, chữ ký)")
    ctx.stats.update(images=1, ocr_lines=len(txt.splitlines()))
    md = f"![Bản gốc]({rel})\n\n## Nội dung chữ\n\n<!-- OCR thô — Claude đối chiếu ảnh để sửa -->\n{txt or '_(OCR không ra chữ)_'}\n"
    return None, md


ROOT_INPUT = None  # docs-input để dò đính kèm trùng (--root)


def convert(src, out, slug, img_prefix="", depth=0):
    """Chuyển 1 file → (ctx, tiêu đề, thân md). Ảnh ghi vào out/images, file đính kèm vào out/source."""
    ext = os.path.splitext(src)[1].lower()
    os.makedirs(out, exist_ok=True)
    ctx = Ctx(src, os.path.abspath(out), slug, img_prefix, depth)

    work = src
    if ext == ".doc":  # Word cũ → docx bằng textutil (có sẵn trên macOS)
        work = os.path.join(ctx.out, "_tmp.docx")
        subprocess.run(["textutil", "-convert", "docx", src, "-output", work], check=True)
        ctx.warn.append("File .doc cũ — đã đổi sang docx bằng textutil, có thể mất hình vẽ; xem bản gốc nếu nghi ngờ")
    ctx.src = work
    kind = {".docx": conv_docx, ".doc": conv_docx, ".pdf": conv_pdf, ".xlsx": conv_xlsx, ".xlsm": conv_xlsx,
            ".xls": conv_xls, ".csv": conv_csv, ".pptx": conv_pptx}.get(ext)
    if kind is None and ext in IMG_EXT:
        kind = conv_image
    if kind is None:
        print(f"Chưa hỗ trợ định dạng {ext}", file=sys.stderr)
        sys.exit(2)
    title, body = kind(ctx)
    if work != src and os.path.exists(work):
        os.remove(work)

    # file đính kèm: chuyển từng file, nối vào cuối tài liệu (đủ nội dung trong 1 .md)
    if ctx.attachments:
        parts = ["\n## Đính kèm\n"]
        for n, (rel, label) in enumerate(ctx.attachments, 1):
            sub = os.path.join(ctx.out, f"_att{n:02d}")
            try:
                sctx, stitle, sbody = convert(os.path.join(ctx.out, rel), sub, f"{slug}-dinh-kem-{n:02d}",
                                              f"{img_prefix}dk{n:02d}-", depth + 1)
            except SystemExit:
                parts.append(f"### Đính kèm {n:02d} — [{label}]({rel})\n\n_(định dạng chưa hỗ trợ — mở file gốc)_\n")
                continue
            for fn in os.listdir(os.path.join(sub, "images")) if os.path.isdir(os.path.join(sub, "images")) else []:
                shutil.move(os.path.join(sub, "images", fn), os.path.join(ctx.img_dir, fn))
            for fn in os.listdir(os.path.join(sub, "source")) if os.path.isdir(os.path.join(sub, "source")) else []:
                shutil.move(os.path.join(sub, "source", fn), os.path.join(ctx.out, "source", fn))
            shutil.rmtree(sub)
            name = stitle or label
            dup = find_duplicate(os.path.join(ctx.out, rel), sbody, ROOT_INPUT)
            if dup:
                ctx.warn.append(f"[Đính kèm {n:02d}] {dup[1]} với `{dup[0]}` — vẫn chép đủ nội dung, ghi chú Trùng với")
            parts.append(f"### Đính kèm {n:02d} — {name}\n\n> File: [{os.path.basename(rel)}]({rel})"
                         + (f" (tên gốc `{label}`)" if label and not label.startswith("Đính kèm") else "")
                         + (f"\n> Trùng với `{dup[0]}` ({dup[1]})" if dup else "") + "\n\n"
                         + fix_md(sbody, 2) + "\n")
            ctx.raw += sctx.raw
            ctx.raw_ocr += sctx.raw_ocr
            ctx.warn += [f"[Đính kèm {n:02d}] {w}" for w in sctx.warn]
            ctx.stats.setdefault("attachments", []).append(dict(file=rel, title=name, duplicate_of=dup, **sctx.stats))
        body = body.rstrip() + "\n" + "\n".join(parts)
    return ctx, title, body


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("--out", required=True)
    ap.add_argument("--slug", required=True)
    ap.add_argument("--date", default=dt.date.today().isoformat(), help="ngày nhận")
    ap.add_argument("--root", default="docs-input", help="thư mục nguồn để dò đính kèm trùng")
    a = ap.parse_args()
    src = os.path.abspath(a.src)
    ext = os.path.splitext(src)[1].lower()
    global ROOT_INPUT
    ROOT_INPUT = os.path.abspath(a.root) if os.path.isdir(a.root) else None
    ctx, title, body = convert(src, os.path.abspath(a.out), a.slug)

    # bản sao gốc
    os.makedirs(os.path.join(ctx.out, "source"), exist_ok=True)
    src_rel = f"source/{a.slug}{ext}"
    shutil.copy2(src, os.path.join(ctx.out, src_rel))
    if not os.listdir(ctx.img_dir):
        os.rmdir(ctx.img_dir)

    orig = nfc(os.path.basename(src))
    title = nfc(title).strip() if title and title.strip() else os.path.splitext(orig)[0]
    kind_vi = {".docx": "Word", ".doc": "Word cũ", ".pdf": "PDF", ".xlsx": "Excel", ".xlsm": "Excel",
               ".xls": "Excel cũ", ".csv": "CSV", ".pptx": "PowerPoint"}.get(ext, "Ảnh")
    head = (f"# {title}\n\n"
            f"> Nguồn: [`{orig}`]({src_rel}) ({kind_vi}) · Nhận: {a.date} · Chuyển sang Markdown bằng `/intake`.\n\n")
    with open(os.path.join(ctx.out, f"{a.slug}.md"), "w", encoding="utf-8") as f:
        f.write(head + nfc(body).strip() + "\n")
    with open(os.path.join(ctx.out, "_raw.txt"), "w", encoding="utf-8") as f:
        f.write(nfc("\n".join(ctx.raw)))
    if ctx.raw_ocr:
        with open(os.path.join(ctx.out, "_raw_ocr.txt"), "w", encoding="utf-8") as f:
            f.write(nfc("\n".join(ctx.raw_ocr)))
    warns = list(dict.fromkeys(ctx.warn))
    rep = dict(source=orig, kind=kind_vi, title=title, md=f"{a.slug}.md", stats=ctx.stats, warnings=warns)
    with open(os.path.join(ctx.out, "_report.json"), "w", encoding="utf-8") as f:
        json.dump(rep, f, ensure_ascii=False, indent=1)
    print(json.dumps(rep, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
